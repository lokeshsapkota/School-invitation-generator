from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from copy import deepcopy
from datetime import datetime

def read_template():
    doc = Document('template.docx')
    return doc

def fill_template(doc, school_name, school_address):
    date = datetime.now().strftime("%Y/%m/%d")

    filled_doc = deepcopy(doc)  # Create a deep copy of the template

    for paragraph in filled_doc.paragraphs:
        for run in paragraph.runs:
            # Check for placeholders and replace them
            if '[School/college]' in run.text:
                run.text = run.text.replace('[School/college]', school_name)
            if '[School/College Name]' in run.text:
                run.text = run.text.replace('[School/College Name]', school_name)
            if '[School/College Address]' in run.text:
                run.text = run.text.replace('[School/College Address]', school_address)
            if '[Date]' in run.text:
                run.text = run.text.replace('[Date]', date)

    return filled_doc

def save_invite(doc, school_name):
    filename = f'{school_name}_invitation.docx'
    doc.save(filename)
    print(f'Invitation saved as {filename}')

def main():
    template = read_template()

    with open('details.txt', 'r') as details_file:
        schools = details_file.readlines()

    for school_info in schools:
        parts = school_info.split(',', 1)
        school_name = parts[0].strip()
        school_address = parts[1].strip() if len(parts) > 1 else ""

        # Fill in the template with school-specific details
        filled_doc = fill_template(template, school_name, school_address)

        # Save the filled document
        save_invite(filled_doc, school_name)

if __name__ == "__main__":
    main()
